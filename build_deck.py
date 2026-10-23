import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette matching Namma Santhe Pitch Deck PDF
    BG_CREAM = RGBColor(250, 246, 235)         # #FAF6EB
    TOP_BAR_ORANGE = RGBColor(232, 106, 51)     # #E86A33
    DARK_GREEN = RGBColor(30, 83, 54)           # #1E5336
    ACCENT_ORANGE = RGBColor(232, 106, 51)      # #E86A33
    CARD_PALE_GREEN = RGBColor(217, 232, 214)   # #D9E8D6
    CARD_PALE_PEACH = RGBColor(244, 221, 194)   # #F4DDC2
    CARD_DARK_GREEN = RGBColor(27, 94, 60)      # #1B5E3C
    TEXT_DARK = RGBColor(35, 47, 40)            # #232F28
    TEXT_MUTED = RGBColor(100, 110, 105)        # #646E69
    WHITE = RGBColor(255, 255, 255)
    LINE_YELLOW = RGBColor(225, 175, 75)

    FONT_FAMILY = "Segoe UI"
    FONT_KANNADA = "Nirmala UI"

    def set_slide_bg(slide):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_CREAM
        bg_shape.line.fill.background()
        return bg_shape

    def add_top_bar(slide):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.24))
        bar.fill.solid()
        bar.fill.fore_color.rgb = TOP_BAR_ORANGE
        bar.line.fill.background()
        return bar

    def add_header(slide, title_text, subtitle_text=None):
        add_top_bar(slide)
        tx_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.48), Inches(11.5), Inches(1.3))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = FONT_FAMILY
            p2.font.size = Pt(17)
            p2.font.bold = True
            p2.font.color.rgb = ACCENT_ORANGE
            p2.space_before = Pt(5)

    # ==========================================
    # SLIDE 1: Title Slide (Cover)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)

    # Hero Artwork on right (Cropped pure art)
    img1_path = "pdf_slides/art_slide1.png"
    if os.path.exists(img1_path):
        s1.shapes.add_picture(img1_path, Inches(8.055), Inches(0), Inches(5.278), Inches(7.5))

    # Left content box
    tbox = s1.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(6.8), Inches(4.8))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p = tf1.paragraphs[0]
    p.text = "Namma Santhe"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN

    p_kn = tf1.add_paragraph()
    p_kn.text = "ನಮ್ಮ ಸಂತೆ"
    p_kn.font.name = FONT_KANNADA
    p_kn.font.size = Pt(38)
    p_kn.font.bold = True
    p_kn.font.color.rgb = DARK_GREEN
    p_kn.space_before = Pt(6)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Local produce, local shops, local love"
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(21)
    p_sub.font.bold = False
    p_sub.font.color.rgb = ACCENT_ORANGE
    p_sub.space_before = Pt(18)

    p_team = tf1.add_paragraph()
    p_team.text = "HEXAPLUSE (TEAM)"
    p_team.font.name = FONT_FAMILY
    p_team.font.size = Pt(14)
    p_team.font.color.rgb = TEXT_MUTED
    p_team.space_before = Pt(120)

    # ==========================================
    # SLIDE 2: The Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "The Problem", "Local products deserve a bigger stage!")

    # Card 1: SELLERS (Peach)
    c1 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.1), Inches(5.5), Inches(4.6))
    c1.fill.solid(); c1.fill.fore_color.rgb = CARD_PALE_PEACH; c1.line.fill.background()
    tf = c1.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.35); tf.margin_top = Inches(0.35); tf.margin_right = Inches(0.35)
    
    p = tf.paragraphs[0]
    p.text = "SELLERS"
    p.font.name = FONT_FAMILY; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = TEXT_DARK
    
    seller_points = [
        "Limited online visibility and no searchable digital storefront",
        "Weak logistics and delivery reach outside walking distance",
        "High commissions (18% – 35% on large aggregators)",
        "Inventory and payment struggles with manual paper records"
    ]
    for pt in seller_points:
        p_pt = tf.add_paragraph()
        p_pt.text = f"• {pt}"
        p_pt.font.name = FONT_FAMILY; p_pt.font.size = Pt(14.5); p_pt.font.color.rgb = TEXT_DARK
        p_pt.space_before = Pt(14)

    # Card 2: CUSTOMERS (Sage Green)
    c2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.9), Inches(2.1), Inches(5.5), Inches(4.6))
    c2.fill.solid(); c2.fill.fore_color.rgb = CARD_PALE_GREEN; c2.line.fill.background()
    tf = c2.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.35); tf.margin_top = Inches(0.35); tf.margin_right = Inches(0.35)
    
    p = tf.paragraphs[0]
    p.text = "CUSTOMERS"
    p.font.name = FONT_FAMILY; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = DARK_GREEN
    
    cust_points = [
        "Few trusted local sellers available online in one place",
        "Fragmented shopping experience across physical markets",
        "Limited product assortment information and unknown stock",
        "Hard to compare nearby options and prices before traveling"
    ]
    for pt in cust_points:
        p_pt = tf.add_paragraph()
        p_pt.text = f"• {pt}"
        p_pt.font.name = FONT_FAMILY; p_pt.font.size = Pt(14.5); p_pt.font.color.rgb = TEXT_DARK
        p_pt.space_before = Pt(14)

    # ==========================================
    # SLIDE 3: Our Solution
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "Our Solution")

    # Center Hub (Dark Green)
    hub = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.8), Inches(2.8), Inches(3.7), Inches(2.2))
    hub.fill.solid(); hub.fill.fore_color.rgb = CARD_DARK_GREEN; hub.line.fill.background()
    tf = hub.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "NAMMA SANTHE\nನಮ್ಮ ಸಂತೆ"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = WHITE

    # Connector lines
    l1 = s3.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.8), Inches(3.5), Inches(3.6), Inches(2.4))
    l1.line.color.rgb = ACCENT_ORANGE; l1.line.width = Pt(2.5)

    l2 = s3.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.8), Inches(4.3), Inches(3.6), Inches(5.4))
    l2.line.color.rgb = ACCENT_ORANGE; l2.line.width = Pt(2.5)

    l3 = s3.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.5), Inches(3.5), Inches(9.7), Inches(2.4))
    l3.line.color.rgb = ACCENT_ORANGE; l3.line.width = Pt(2.5)

    l4 = s3.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.5), Inches(4.3), Inches(9.7), Inches(5.4))
    l4.line.color.rgb = ACCENT_ORANGE; l4.line.width = Pt(2.5)

    # Satellites
    s_tl = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(1.9), Inches(2.8), Inches(1.4))
    s_tl.fill.solid(); s_tl.fill.fore_color.rgb = CARD_PALE_GREEN; s_tl.line.fill.background()
    tf = s_tl.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Digital storefronts\n(Kannada & English Catalog)"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(14.5); p.font.color.rgb = TEXT_DARK; p.font.bold = True

    s_bl = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(4.9), Inches(2.8), Inches(1.4))
    s_bl.fill.solid(); s_bl.fill.fore_color.rgb = CARD_PALE_GREEN; s_bl.line.fill.background()
    tf = s_bl.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Secure payment\n(Instant UPI & Escrow)"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(14.5); p.font.color.rgb = TEXT_DARK; p.font.bold = True

    s_tr = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.2), Inches(1.9), Inches(2.8), Inches(1.4))
    s_tr.fill.solid(); s_tr.fill.fore_color.rgb = CARD_PALE_PEACH; s_tr.line.fill.background()
    tf = s_tr.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Orders\n(Real-time live seller alerts)"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(14.5); p.font.color.rgb = TEXT_DARK; p.font.bold = True

    s_br = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.2), Inches(4.9), Inches(2.8), Inches(1.4))
    s_br.fill.solid(); s_br.fill.fore_color.rgb = CARD_PALE_PEACH; s_br.line.fill.background()
    tf = s_br.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Delivery tracking\n(Hyperlocal partner routing)"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(14.5); p.font.color.rgb = TEXT_DARK; p.font.bold = True

    tag_box = s3.shapes.add_textbox(Inches(3.0), Inches(6.6), Inches(7.333), Inches(0.6))
    p = tag_box.text_frame.paragraphs[0]; p.text = "Seamless. Simple. Local."; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ACCENT_ORANGE

    # ==========================================
    # SLIDE 4: How It Works
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "How It Works")

    steps = [
        ("1 Discover", "Browse nearby verified santhes & neighborhood shops"),
        ("2 Compare", "Check live rates for fresh produce across local stalls"),
        ("3 Order", "Add fresh items to cart with single-click ordering"),
        ("4 Pay", "Fast, secure UPI QR, card, or Cash on Delivery"),
        ("5 Track", "Real-time rider updates right to your doorstep"),
        ("6 Support Local", "Direct fair payment to local farmers and shopkeepers")
    ]
    
    top_pos = 1.9
    for i, (title, desc) in enumerate(steps):
        tag = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(top_pos), Inches(0.12), Inches(0.6))
        tag.fill.solid(); tag.fill.fore_color.rgb = ACCENT_ORANGE if i % 2 == 0 else LINE_YELLOW; tag.line.fill.background()

        st_box = s4.shapes.add_textbox(Inches(1.15), Inches(top_pos - 0.05), Inches(5.3), Inches(0.75))
        tf = st_box.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.text = title
        p.font.name = FONT_FAMILY; p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph(); p2.text = desc
        p2.font.name = FONT_FAMILY; p2.font.size = Pt(12); p2.font.color.rgb = TEXT_MUTED

        top_pos += 0.85

    # Cropped artwork on right
    img4_path = "pdf_slides/art_slide4.png"
    if os.path.exists(img4_path):
        s4.shapes.add_picture(img4_path, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))

    # ==========================================
    # SLIDE 5: Key Features
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "Key Features")

    features = [
        ("Nearby discovery", "GPS radius locator, open/close status, and neighborhood santhe clusters.", CARD_PALE_GREEN, DARK_GREEN),
        ("Smart compare", "Real-time price comparison of fresh produce across local sellers.", CARD_PALE_PEACH, ACCENT_ORANGE),
        ("Personal offers", "Festive discounts, bulk vegetable baskets, and community perks.", CARD_PALE_GREEN, DARK_GREEN),
        ("Seller analytics", "Daily revenue, order volume, bestsellers, and peak hours tracking.", CARD_PALE_PEACH, ACCENT_ORANGE),
        ("Digital dukaan", "Kannada & English storefront, instant photo catalog, and stock toggles.", CARD_PALE_GREEN, DARK_GREEN),
        ("Secure wallet", "Direct UPI payouts to sellers with zero predatory platform commissions.", CARD_PALE_PEACH, ACCENT_ORANGE),
    ]

    col_coords = [Inches(0.9), Inches(4.9), Inches(8.9)]
    row_coords = [Inches(2.0), Inches(4.6)]

    for idx, (title, desc, bg_color, dot_color) in enumerate(features):
        c_idx = idx % 3
        r_idx = idx // 3
        
        card = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, col_coords[c_idx], row_coords[r_idx], Inches(3.6), Inches(2.2))
        card.fill.solid(); card.fill.fore_color.rgb = bg_color; card.line.fill.background()
        
        dot = s5.shapes.add_shape(MSO_SHAPE.OVAL, col_coords[c_idx] + Inches(1.5), row_coords[r_idx] + Inches(0.25), Inches(0.6), Inches(0.6))
        dot.fill.solid(); dot.fill.fore_color.rgb = dot_color; dot.line.fill.background()

        tbox = s5.shapes.add_textbox(col_coords[c_idx] + Inches(0.2), row_coords[r_idx] + Inches(0.95), Inches(3.2), Inches(1.1))
        tf = tbox.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]; p.text = title; p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_FAMILY; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph(); p2.text = desc; p2.alignment = PP_ALIGN.CENTER
        p2.font.name = FONT_FAMILY; p2.font.size = Pt(11.5); p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

    # ==========================================
    # SLIDE 6: User Journey
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "User Journey")

    journey_steps = [
        ("1", "Browse local shops", CARD_DARK_GREEN),
        ("2", "Add to cart", CARD_DARK_GREEN),
        ("3", "Checkout", CARD_DARK_GREEN),
        ("4", "Delivery", ACCENT_ORANGE),
    ]

    cur_top = 2.1
    for num, txt, color in journey_steps:
        pill = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(cur_top), Inches(0.55), Inches(0.9))
        pill.fill.solid(); pill.fill.fore_color.rgb = color; pill.line.fill.background()
        tf = pill.text_frame
        p = tf.paragraphs[0]; p.text = num; p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = WHITE

        t_box = s6.shapes.add_textbox(Inches(1.65), Inches(cur_top + 0.22), Inches(4.5), Inches(0.6))
        tf = t_box.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.text = txt
        p.font.name = FONT_FAMILY; p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = TEXT_DARK

        cur_top += 1.05

    motto_box = s6.shapes.add_textbox(Inches(0.9), Inches(6.3), Inches(5.5), Inches(0.9))
    tf = motto_box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.text = "Shop Local. Support Small."
    p.font.name = FONT_FAMILY; p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = ACCENT_ORANGE
    p2 = tf.add_paragraph(); p2.text = "Build Stronger Communities."
    p2.font.name = FONT_FAMILY; p2.font.size = Pt(16.5); p2.font.bold = True; p2.font.color.rgb = ACCENT_ORANGE

    # Cropped artwork on right
    img6_path = "pdf_slides/art_slide6.png"
    if os.path.exists(img6_path):
        s6.shapes.add_picture(img6_path, Inches(7.0), Inches(1.8), Inches(5.5), Inches(5.0))

    # ==========================================
    # SLIDE 7: Seller Dashboard
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "Seller Dashboard", "Empowering local micro-merchants with actionable growth metrics")

    metrics = [
        ("₹1.05L", "Revenue", CARD_PALE_GREEN, DARK_GREEN),
        ("82", "Orders", CARD_PALE_GREEN, DARK_GREEN),
        ("4", "Stock alerts", CARD_PALE_PEACH, ACCENT_ORANGE),
    ]

    m_top = 2.2
    for val, lbl, bg, fg in metrics:
        card = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(m_top), Inches(3.4), Inches(1.3))
        card.fill.solid(); card.fill.fore_color.rgb = bg; card.line.fill.background()
        tf = card.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.3); tf.margin_top = Inches(0.2)
        
        p = tf.paragraphs[0]; p.text = val
        p.font.name = FONT_FAMILY; p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = fg
        
        p2 = tf.add_paragraph(); p2.text = lbl
        p2.font.name = FONT_FAMILY; p2.font.size = Pt(13); p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(2)

        m_top += 1.6

    # Native Line Chart on Right
    chart_data = CategoryChartData()
    chart_data.categories = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun']
    chart_data.add_series('Sales ₹000', (42, 54, 61, 74, 88, 105))
    chart_data.add_series('Orders', (28, 34, 48, 57, 70, 82))

    ctitle = s7.shapes.add_textbox(Inches(4.9), Inches(2.1), Inches(7.5), Inches(0.6))
    p = ctitle.text_frame.paragraphs[0]
    p.text = "ಮಾಸಿಕ ಬೆಳವಣಿಗೆ | Monthly growth"
    p.font.name = FONT_KANNADA; p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = DARK_GREEN

    x, y, cx, cy = Inches(4.7), Inches(2.7), Inches(7.8), Inches(4.2)
    chart = s7.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x, y, cx, cy, chart_data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False

    # Style series colors matching brand
    series1 = chart.series[0]
    series1.format.line.color.rgb = DARK_GREEN
    series1.format.line.width = Pt(3)

    series2 = chart.series[1]
    series2.format.line.color.rgb = ACCENT_ORANGE
    series2.format.line.width = Pt(3)

    # ==========================================
    # SLIDE 8: Customer View
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "Customer View", "A familiar local shopping experience — now in your pocket.")

    t_box = s8.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(5.8), Inches(4.5))
    tf = t_box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    c_points = [
        "Nearby shop listings and authentic community ratings",
        "Kannada & English bilingual product descriptions",
        "Clear product details, transparent weight units, and live prices",
        "Simple 1-tap cart and secure checkout with UPI",
        "Local vendor identity and farm stories front and center"
    ]
    p = tf.paragraphs[0]
    p.text = f"• {c_points[0]}"
    p.font.name = FONT_FAMILY; p.font.size = Pt(16); p.font.color.rgb = TEXT_DARK
    
    for pt in c_points[1:]:
        p_pt = tf.add_paragraph()
        p_pt.text = f"• {pt}"
        p_pt.font.name = FONT_FAMILY; p_pt.font.size = Pt(16); p_pt.font.color.rgb = TEXT_DARK
        p_pt.space_before = Pt(16)

    # Cropped App Mockup Frame
    img8_path = "pdf_slides/art_slide8.png"
    if os.path.exists(img8_path):
        s8.shapes.add_picture(img8_path, Inches(7.5), Inches(1.6), Inches(4.5), Inches(5.4))

    # ==========================================
    # SLIDE 9: Application Architecture & Database Design
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "Application Architecture & Database Design")

    # Layer 1: Client
    lay1 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.0), Inches(7.8), Inches(1.25))
    lay1.fill.solid(); lay1.fill.fore_color.rgb = CARD_PALE_GREEN; lay1.line.fill.background()
    tf = lay1.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "CUSTOMER PWA • SELLER PORTAL"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = DARK_GREEN
    p2 = tf.add_paragraph(); p2.text = "React 19 • TypeScript • Responsive Viewport • Tailwind CSS • Vite"; p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_FAMILY; p2.font.size = Pt(12); p2.font.color.rgb = TEXT_MUTED

    a1 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.7), Inches(3.25), Inches(0.08), Inches(0.35))
    a1.fill.solid(); a1.fill.fore_color.rgb = ACCENT_ORANGE; a1.line.fill.background()

    # Layer 2: API & Services
    lay2 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.6), Inches(7.8), Inches(1.3))
    lay2.fill.solid(); lay2.fill.fore_color.rgb = CARD_PALE_PEACH; lay2.line.fill.background()
    tf = lay2.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "API & BACKEND SERVICES"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_ORANGE
    p2 = tf.add_paragraph(); p2.text = "Product Catalog • Live Orders • Geo-Distance Engine • Merchant Analytics"; p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_FAMILY; p2.font.size = Pt(12); p2.font.color.rgb = TEXT_MUTED

    a2 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.7), Inches(4.9), Inches(0.08), Inches(0.35))
    a2.fill.solid(); a2.fill.fore_color.rgb = ACCENT_ORANGE; a2.line.fill.background()

    # Layer 3: Database
    lay3 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.25), Inches(7.8), Inches(1.25))
    lay3.fill.solid(); lay3.fill.fore_color.rgb = CARD_PALE_GREEN; lay3.line.fill.background()
    tf = lay3.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "DATABASE & STORAGE"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = DARK_GREEN
    p2 = tf.add_paragraph(); p2.text = "Categories • Products • Orders & Line Items • Sellers • Inventory Locks"; p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_FAMILY; p2.font.size = Pt(12); p2.font.color.rgb = TEXT_MUTED

    # Right Integrations
    lp = s9.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.7), Inches(4.2), Inches(9.8), Inches(2.9))
    lp.line.color.rgb = DARK_GREEN; lp.line.width = Pt(2.5)

    ld = s9.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.7), Inches(4.2), Inches(9.8), Inches(5.3))
    ld.line.color.rgb = ACCENT_ORANGE; ld.line.width = Pt(2.5)

    card_pay = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.6), Inches(2.3), Inches(2.8), Inches(1.5))
    card_pay.fill.solid(); card_pay.fill.fore_color.rgb = CARD_DARK_GREEN; card_pay.line.fill.background()
    tf = card_pay.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Payment Gateway\nUPI 2.0 • QR • Razorpay"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(14.5); p.font.bold = True; p.font.color.rgb = WHITE

    card_del = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.6), Inches(4.7), Inches(2.8), Inches(1.5))
    card_del.fill.solid(); card_del.fill.fore_color.rgb = ACCENT_ORANGE; card_del.line.fill.background()
    tf = card_del.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Delivery Services\nHyperlocal Rider Routing"; p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(14.5); p.font.bold = True; p.font.color.rgb = WHITE

    bot_box = s9.shapes.add_textbox(Inches(3.0), Inches(6.8), Inches(7.333), Inches(0.5))
    p = bot_box.text_frame.paragraphs[0]
    p.text = "Scalable • Robust • Secure"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = DARK_GREEN

    # ==========================================
    # SLIDE 10: Security & Inventory Reliability
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "Security & Inventory Reliability")

    items = [
        ("Secure payments", CARD_DARK_GREEN),
        ("Verified sellers", ACCENT_ORANGE),
        ("Real-time stock", CARD_DARK_GREEN),
        ("Quality assurance", ACCENT_ORANGE),
        ("Customer support", CARD_DARK_GREEN),
    ]

    p_top = 2.1
    for label, color in items:
        pill = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(p_top), Inches(0.55), Inches(0.7))
        pill.fill.solid(); pill.fill.fore_color.rgb = color; pill.line.fill.background()
        tf = pill.text_frame
        p = tf.paragraphs[0]; p.text = "✓"; p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_FAMILY; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = WHITE

        lbox = s10.shapes.add_textbox(Inches(1.65), Inches(p_top + 0.12), Inches(5.2), Inches(0.5))
        tf = lbox.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.text = label
        p.font.name = FONT_FAMILY; p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = TEXT_DARK

        p_top += 0.95

    rblock = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.6), Inches(2.1), Inches(4.8), Inches(4.4))
    rblock.fill.solid(); rblock.fill.fore_color.rgb = CARD_DARK_GREEN; rblock.line.fill.background()
    tf = rblock.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.5); tf.margin_right = Inches(0.5)
    
    p = tf.paragraphs[0]
    p.text = "Accurate stock\n\nProtected payments\n\nResponsive support\n\n100% Trust Guarantee"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY; p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 11: Development Roadmap
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11)
    add_header(s11, "Development Roadmap", "Execution milestones from MVP to citywide community scale")

    track = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(2.6), Inches(10.3), Inches(0.08))
    track.fill.solid(); track.fill.fore_color.rgb = LINE_YELLOW; track.line.fill.background()

    phases = [
        ("1", "01", "Research & MVP", ["Hyperlocal catalog", "Customer PWA & cart", "Pilot vendor testing"], CARD_PALE_GREEN, CARD_DARK_GREEN),
        ("2", "02", "Seller onboarding", ["Vernacular listing", "Direct UPI payout", "Merchant inventory app"], CARD_PALE_PEACH, ACCENT_ORANGE),
        ("3", "03", "Enhance & scale", ["AI price intelligence", "Batch route dispatch", "Multi-santhe clusters"], CARD_PALE_GREEN, CARD_DARK_GREEN),
        ("4", "04", "Market & engage", ["ONDC integration", "Farmer collectives", "Citywide rollout"], CARD_PALE_PEACH, ACCENT_ORANGE),
    ]

    col_positions = [Inches(0.9), Inches(3.9), Inches(6.9), Inches(9.9)]

    for i, (num_str, phase_id, title, pts, card_bg, badge_color) in enumerate(phases):
        cp = col_positions[i]
        
        circle = s11.shapes.add_shape(MSO_SHAPE.OVAL, cp + Inches(0.85), Inches(2.2), Inches(0.85), Inches(0.85))
        circle.fill.solid(); circle.fill.fore_color.rgb = badge_color; circle.line.fill.background()
        tf = circle.text_frame
        p = tf.paragraphs[0]; p.text = num_str; p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_FAMILY; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = WHITE

        box = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, cp, Inches(3.5), Inches(2.6), Inches(3.0))
        box.fill.solid(); box.fill.fore_color.rgb = card_bg; box.line.fill.background()
        tf = box.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.25); tf.margin_right = Inches(0.2)
        
        p = tf.paragraphs[0]; p.text = phase_id; p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_FAMILY; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = TEXT_MUTED

        p2 = tf.add_paragraph(); p2.text = title; p2.alignment = PP_ALIGN.CENTER
        p2.font.name = FONT_FAMILY; p2.font.size = Pt(15); p2.font.bold = True; p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(4)

        for pt in pts:
            p_pt = tf.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.name = FONT_FAMILY; p_pt.font.size = Pt(11); p_pt.font.color.rgb = TEXT_DARK
            p_pt.space_before = Pt(6)

    # ==========================================
    # SLIDE 12: AI Tools & Tech Stack (Plan, Prototype & Build)
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12)
    add_header(s12, "AI Tools & Tech Stack", "Plan, prototype & build")

    stack_1 = [
        ("REQUIREMENTS, IDEATION & BRAINSTORMING", "Google AI Studio, Gemini, ChatGPT, Claude",
         "Generates PRDs, database schemas, and system architectures. Gathers requirements, outlines business logic, user roles, and hyperlocal discovery features.", CARD_PALE_GREEN),
        ("PROTOTYPING & CODING", "Cursor, v0 by Vercel & Lovable",
         "Rapidly scaffolds responsive UI components and writes full-stack code. Instantly turns natural language prompts into working user interfaces.", CARD_PALE_PEACH),
        ("AGENTIC DEVELOPMENT", "Antigravity & Open Code",
         "Agentic development platforms where AI agents plan, build, run, and test application features and integrations autonomously.", CARD_PALE_PEACH),
        ("ANALYTICS & INTELLIGENCE", "Gemini API",
         "Powers smart personalized product recommendations, vernacular voice search understanding, and bilingual Kannada-English catalog translation.", CARD_PALE_GREEN),
    ]

    coords_2x2 = [
        (Inches(0.9), Inches(2.1)),
        (Inches(6.9), Inches(2.1)),
        (Inches(0.9), Inches(4.5)),
        (Inches(6.9), Inches(4.5)),
    ]

    for (cat, tools, desc, bg), (x, y) in zip(stack_1, coords_2x2):
        card = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(5.5), Inches(2.15))
        card.fill.solid(); card.fill.fore_color.rgb = bg; card.line.fill.background()
        tf = card.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.3); tf.margin_top = Inches(0.2); tf.margin_right = Inches(0.3)
        
        p = tf.paragraphs[0]; p.text = cat
        p.font.name = FONT_FAMILY; p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = ACCENT_ORANGE

        p2 = tf.add_paragraph(); p2.text = tools
        p2.font.name = FONT_FAMILY; p2.font.size = Pt(14); p2.font.bold = True; p2.font.color.rgb = DARK_GREEN
        p2.space_before = Pt(3)

        p3 = tf.add_paragraph(); p3.text = desc
        p3.font.name = FONT_FAMILY; p3.font.size = Pt(11.5); p3.font.color.rgb = TEXT_DARK
        p3.space_before = Pt(4)

    # ==========================================
    # SLIDE 13: AI Tools & Tech Stack (Test, Document, Design & Ship)
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s13)
    add_header(s13, "AI Tools & Tech Stack", "Test, document, design & ship")

    stack_2 = [
        ("TESTING & DEPLOYMENT", "CodiumAI, Oxlint & Vercel",
         "Automatically generates unit tests, runs rapid static code analysis, and streamlines continuous cloud deployment with preview URLs.", CARD_PALE_PEACH),
        ("RESEARCH & DOCUMENTATION", "Obsidian & Gravity Write",
         "Obsidian organizes hyper-local research notes, market studies, and architectural docs. Gravity Write generates crisp customer-facing copy.", CARD_PALE_GREEN),
        ("DESIGN & VISUAL ASSETS", "Canva & Generative AI Studio",
         "Designs branding identities, logos, santhe market illustrations, and high-fidelity promotional banners.", CARD_PALE_GREEN),
        ("MANAGEMENT & VERSION CONTROL", "GitHub",
         "Handles source code tracking, branch protection rules, issue prioritization, and automated pairing workflows.", CARD_PALE_PEACH),
    ]

    for (cat, tools, desc, bg), (x, y) in zip(stack_2, coords_2x2):
        card = s13.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(5.5), Inches(2.15))
        card.fill.solid(); card.fill.fore_color.rgb = bg; card.line.fill.background()
        tf = card.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.3); tf.margin_top = Inches(0.2); tf.margin_right = Inches(0.3)
        
        p = tf.paragraphs[0]; p.text = cat
        p.font.name = FONT_FAMILY; p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = ACCENT_ORANGE

        p2 = tf.add_paragraph(); p2.text = tools
        p2.font.name = FONT_FAMILY; p2.font.size = Pt(14); p2.font.bold = True; p2.font.color.rgb = DARK_GREEN
        p2.space_before = Pt(3)

        p3 = tf.add_paragraph(); p3.text = desc
        p3.font.name = FONT_FAMILY; p3.font.size = Pt(11.5); p3.font.color.rgb = TEXT_DARK
        p3.space_before = Pt(4)

    # ==========================================
    # SLIDE 14: Closing Slide
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s14)

    # Artwork on right (Cropped pure art)
    img14_path = "pdf_slides/art_slide14.png"
    if os.path.exists(img14_path):
        s14.shapes.add_picture(img14_path, Inches(8.19), Inches(0), Inches(5.14), Inches(7.5))

    # Left content
    cbox = s14.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(6.8), Inches(5.2))
    tf14 = cbox.text_frame; tf14.word_wrap = True
    tf14.margin_left = tf14.margin_top = tf14.margin_right = tf14.margin_bottom = 0

    p = tf14.paragraphs[0]
    p.text = "Support Local. Shop Local.\nGrow Local."
    p.font.name = FONT_FAMILY; p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = DARK_GREEN

    p_kn = tf14.add_paragraph()
    p_kn.text = "ಸ್ಥಳೀಯವನ್ನು ಬೆಂಬಲಿಸಿ. ಸ್ಥಳೀಯವಾಗಿ ಖರೀದಿಸಿ.\nಒಟ್ಟಾಗಿ ಬೆಳೆಯೋಣ."
    p_kn.font.name = FONT_KANNADA; p_kn.font.size = Pt(22); p_kn.font.bold = True; p_kn.font.color.rgb = ACCENT_ORANGE
    p_kn.space_before = Pt(16)

    p_pil = tf14.add_paragraph()
    p_pil.text = "Support local businesses • Increase farmer incomes\nReduce food waste • Build community resilience"
    p_pil.font.name = FONT_FAMILY; p_pil.font.size = Pt(14); p_pil.font.color.rgb = TEXT_MUTED
    p_pil.space_before = Pt(28)

    p_ty = tf14.add_paragraph()
    p_ty.text = "| Thank you"
    p_ty.font.name = FONT_FAMILY; p_ty.font.size = Pt(22); p_ty.font.bold = True; p_ty.font.color.rgb = DARK_GREEN
    p_ty.space_before = Pt(45)

    p_hex = tf14.add_paragraph()
    p_hex.text = "TEAM HEXAPLUSE"
    p_hex.font.name = FONT_FAMILY; p_hex.font.size = Pt(14); p_hex.font.bold = True; p_hex.font.color.rgb = TEXT_MUTED
    p_hex.space_before = Pt(6)

    # Save Presentation
    output_filename = "Namma_Santhe_Master_Pitch_Deck.pptx"
    prs.save(output_filename)
    print(f"Successfully generated {output_filename} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    create_presentation()
